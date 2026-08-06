"""Document parser: extracts text and structure from PDF, DOCX, PPTX, TXT."""
from __future__ import annotations
import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional
from app.utils.logger import get_logger
from app.utils.text_cleaner import clean_text, assess_ocr_quality

logger = get_logger(__name__)


@dataclass
class ParsedDocument:
    filename: str
    file_type: str
    full_text: str
    pages: list[str] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)
    page_count: int = 0
    word_count: int = 0

    def __post_init__(self) -> None:
        self.word_count = len(self.full_text.split())


def _extract_with_fitz(path: Path) -> tuple[list[str], dict]:
    """Extract text using PyMuPDF (fast, layout-aware)."""
    import fitz  # PyMuPDF
    doc = fitz.open(str(path))
    raw_meta = doc.metadata or {}
    metadata = {
        "title": raw_meta.get("title", ""),
        "author": raw_meta.get("author", ""),
        "subject": raw_meta.get("subject", ""),
        "creator": raw_meta.get("creator", ""),
    }
    pages = [page.get_text("text") for page in doc]
    doc.close()
    return pages, metadata


def _extract_with_pypdf(path: Path) -> tuple[list[str], dict]:
    """Extract text using pypdf (pure-Python fallback)."""
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    raw_meta = reader.metadata or {}
    metadata = {
        "title": raw_meta.get("/Title", "") or "",
        "author": raw_meta.get("/Author", "") or "",
        "subject": raw_meta.get("/Subject", "") or "",
        "creator": raw_meta.get("/Creator", "") or "",
    }
    pages = [page.extract_text() or "" for page in reader.pages]
    return pages, metadata


def _extract_with_ocr(path: Path) -> tuple[list[str], dict]:
    """Extract text by rendering pages to images + OCR (best for scanned PDFs).

    Requires ``pytesseract`` and a Tesseract binary. Falls back gracefully.
    """
    try:
        import fitz
        import pytesseract  # type: ignore
        from PIL import Image  # type: ignore
    except (ImportError, Exception):
        logger.warning("ocr_unavailable_falling_back")
        return _extract_with_fitz(path)

    doc = fitz.open(str(path))
    pages: list[str] = []
    for page in doc:
        pix = page.get_pixmap(dpi=220)
        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        try:
            text = pytesseract.image_to_string(img)
        except Exception:
            text = ""
        pages.append(text)
    doc.close()
    return pages, {}


def _parse_pdf(path: Path) -> ParsedDocument:
    """Parse a PDF, retrying with a better strategy if quality is poor."""
    strategies = [
        ("fitz", _extract_with_fitz),
        ("pypdf", _extract_with_pypdf),
        ("ocr", _extract_with_ocr),
    ]

    best: Optional[ParsedDocument] = None
    best_score = -1.0

    for label, extractor in strategies:
        try:
            pages, metadata = extractor(path)
        except Exception as e:
            logger.warning(
                "pdf_extract_strategy_failed",
                strategy=label,
                error=str(e),
                filename=path.name,
            )
            continue

        full_text_raw = "\n\n".join(pages)
        full_text = clean_text(full_text_raw)
        quality = assess_ocr_quality(full_text)

        parsed = ParsedDocument(
            filename=path.name,
            file_type="pdf",
            full_text=full_text,
            pages=pages,
            metadata=metadata,
            page_count=len(pages),
        )

        logger.info(
            "pdf_extract_strategy",
            strategy=label,
            filename=path.name,
            score=quality["score"],
            chars=len(full_text),
        )

        if quality["score"] > best_score:
            best_score = quality["score"]
            best = parsed

        # A high-quality extraction is good enough; stop early.
        if quality["score"] >= 0.6:
            break

    if best is None:
        # Final fallback: attempt pypdf directly, else empty doc.
        try:
            pages, metadata = _extract_with_pypdf(path)
            full_text = clean_text("\n\n".join(pages))
            best = ParsedDocument(
                filename=path.name,
                file_type="pdf",
                full_text=full_text,
                pages=pages,
                metadata=metadata,
                page_count=len(pages),
            )
        except Exception as e:
            logger.error("pdf_extract_all_failed", error=str(e), filename=path.name)
            best = ParsedDocument(
                filename=path.name,
                file_type="pdf",
                full_text="",
                pages=[],
                metadata={},
                page_count=0,
            )

    return best


def _parse_docx(path: Path) -> ParsedDocument:
    from docx import Document  # type: ignore
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    full_text = clean_text("\n\n".join(paragraphs))
    return ParsedDocument(
        filename=path.name,
        file_type="docx",
        full_text=full_text,
        pages=[full_text],
        metadata={},
        page_count=1,
    )


def _parse_pptx(path: Path) -> ParsedDocument:
    from pptx import Presentation  # type: ignore
    prs = Presentation(str(path))
    slides = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text)
        slides.append("\n".join(texts))
    full_text = clean_text("\n\n".join(slides))
    return ParsedDocument(
        filename=path.name,
        file_type="pptx",
        full_text=full_text,
        pages=slides,
        metadata={},
        page_count=len(slides),
    )


def _parse_txt(path: Path) -> ParsedDocument:
    text = path.read_text(encoding="utf-8", errors="replace")
    full_text = clean_text(text)
    return ParsedDocument(
        filename=path.name,
        file_type="txt",
        full_text=full_text,
        pages=[full_text],
        metadata={},
        page_count=1,
    )


PARSERS = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".pptx": _parse_pptx,
    ".ppt": _parse_pptx,
    ".txt": _parse_txt,
}


def parse_document(path: Path) -> ParsedDocument:
    """Parse a document file and return structured text content."""
    ext = path.suffix.lower()
    parser = PARSERS.get(ext)
    if not parser:
        raise ValueError(f"Unsupported file type: {ext}. Supported: {list(PARSERS)}")
    logger.info("parsing_document", filename=path.name, file_type=ext)
    doc = parser(path)
    logger.info(
        "document_parsed",
        filename=path.name,
        pages=doc.page_count,
        words=doc.word_count,
    )
    return doc
